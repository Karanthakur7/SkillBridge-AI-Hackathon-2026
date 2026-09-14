import io
import re
from typing import Dict, Any, List


def parse_pdf(file_bytes: bytes) -> str:
    text = ""
    # Try pypdf first
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    except Exception:
        pass

    # If text is empty or very short, fallback to pdfplumber
    if len(text.strip()) < 20:
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
        except Exception:
            pass

    return text.strip()


def parse_docx(file_bytes: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text:
                full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        full_text.append(cell.text)
        return "\n".join(full_text).strip()
    except Exception as e:
        return f"Error extracting DOCX: {str(e)}"


def parse_txt(file_bytes: bytes) -> str:
    for encoding in ["utf-8", "latin-1", "windows-1252", "utf-16"]:
        try:
            return file_bytes.decode(encoding).strip()
        except Exception:
            continue
    return ""


def extract_candidate_name(text: str) -> str:
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    if not lines:
        return "Candidate"

    for line in lines[:5]:
        # Filter out lines that look like emails, urls, or section headers
        if re.search(r"(@|http|www|github|linkedin|resume|curriculum|phone|email|\+?\d{10})", line, re.I):
            continue
        words = line.split()
        if 1 <= len(words) <= 4 and all(w.isalpha() or w in [".", "-", "'"] for w in words):
            return " ".join([w.capitalize() for w in words])
    
    return lines[0][:30] if lines else "Student"


def extract_email(text: str) -> str:
    match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text)
    return match.group(0) if match else ""


def extract_phone(text: str) -> str:
    # Match various phone number formats (including +91, with spaces/dashes)
    match = re.search(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}|\+91[-.\s]?\d{10}|\b\d{10}\b", text)
    return match.group(0).strip() if match else ""


def extract_github(text: str) -> str:
    match = re.search(r"(?:https?:\/\/)?(?:www\.)?github\.com\/([a-zA-Z0-9_\-]+)", text, re.I)
    if match:
        return f"https://github.com/{match.group(1)}"
    match_handle = re.search(r"github:\s*([a-zA-Z0-9_\-]+)", text, re.I)
    if match_handle:
        return f"https://github.com/{match_handle.group(1)}"
    return ""


def extract_linkedin(text: str) -> str:
    match = re.search(r"(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/([a-zA-Z0-9_\-%]+)", text, re.I)
    if match:
        return f"https://linkedin.com/in/{match.group(1)}"
    match_handle = re.search(r"linkedin:\s*([a-zA-Z0-9_\-%]+)", text, re.I)
    if match_handle:
        return f"https://linkedin.com/in/{match_handle.group(1)}"
    return ""


def extract_degree(text: str) -> str:
    degree_patterns = [
        r"(Bachelor of Technology|B\.Tech|BTech|B\.E\.|BE|Bachelor of Engineering)",
        r"(Master of Technology|M\.Tech|MTech|M\.E\.|ME|Master of Engineering)",
        r"(Bachelor of Science|B\.Sc|BSc|B\.C\.A\.|BCA|Bachelor of Computer Applications)",
        r"(Master of Science|M\.Sc|MSc|M\.C\.A\.|MCA|Master of Computer Applications)",
        r"(Bachelor of Business Administration|BBA|Master of Business Administration|MBA)",
        r"(Bachelor of Commerce|B\.Com|BCom|Master of Commerce|M\.Com|MCom)",
        r"(Bachelor of Arts|BA|Master of Arts|MA)"
    ]
    for pattern in degree_patterns:
        match = re.search(pattern, text, re.I)
        if match:
            # Try to grab major if nearby, e.g. "B.Tech in Computer Science"
            submatch = re.search(pattern + r"(?:\s+in\s+([A-Za-z\s]+))?", text, re.I)
            if submatch and submatch.group(2):
                major = submatch.group(2).strip().split("\n")[0][:30]
                return f"{match.group(0)} in {major}"
            return match.group(0)
    return "Bachelor's Degree"


def extract_institution(text: str) -> str:
    patterns = [
        r"([A-Za-z\s]+(?:Institute of Technology|University|College of Engineering|National Institute of Technology|Indian Institute of Technology|State University|Academy|College))",
        r"(IIT\s+[A-Za-z]+|NIT\s+[A-Za-z]+|BITS\s+[A-Za-z]+|IIIT\s+[A-Za-z]+)"
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.I)
        if match:
            clean = match.group(0).strip().split("\n")[0]
            if len(clean) < 60:
                return clean
    return "University / Institute"


def extract_graduation_year(text: str) -> str:
    match = re.search(r"\b(201[5-9]|202[0-9]|2030)\b", text)
    return match.group(0) if match else "2026"


def extract_sections(text: str) -> Dict[str, List[str]]:
    projects = []
    internships = []
    certifications = []

    # Heuristic bullet or line detection under section headers
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    current_section = None

    for line in lines:
        lower = line.lower()
        if re.search(r"^(projects|academic projects|key projects)\b", lower, re.I):
            current_section = "projects"
            continue
        elif re.search(r"^(internships|work experience|experience|internship experience)\b", lower, re.I):
            current_section = "internships"
            continue
        elif re.search(r"^(certifications|certificates|courses|licenses)\b", lower, re.I):
            current_section = "certifications"
            continue
        elif re.search(r"^(skills|education|contact|achievements|hobbies|interests|summary|objective)\b", lower, re.I):
            current_section = None
            continue

        if current_section == "projects" and len(line) > 8:
            if line.startswith(("-", "•", "*")) or re.match(r"^\d+\.", line) or len(projects) < 4:
                clean = re.sub(r"^[-•*\d.]+\s*", "", line)
                if clean and clean not in projects and len(clean) > 5:
                    projects.append(clean[:120])
        elif current_section == "internships" and len(line) > 8:
            if line.startswith(("-", "•", "*")) or re.match(r"^\d+\.", line) or len(internships) < 4:
                clean = re.sub(r"^[-•*\d.]+\s*", "", line)
                if clean and clean not in internships and len(clean) > 5:
                    internships.append(clean[:120])
        elif current_section == "certifications" and len(line) > 6:
            clean = re.sub(r"^[-•*\d.]+\s*", "", line)
            if clean and clean not in certifications and len(clean) > 4:
                certifications.append(clean[:100])

    return {
        "projects": projects[:5],
        "internships": internships[:4],
        "certifications": certifications[:5]
    }


def parse_resume_content(text: str) -> Dict[str, Any]:
    sections = extract_sections(text)
    return {
        "name": extract_candidate_name(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "github_url": extract_github(text),
        "linkedin_url": extract_linkedin(text),
        "degree": extract_degree(text),
        "institution": extract_institution(text),
        "graduation_year": extract_graduation_year(text),
        "projects": sections["projects"],
        "internships": sections["internships"],
        "certifications": sections["certifications"],
        "raw_text": text
    }
